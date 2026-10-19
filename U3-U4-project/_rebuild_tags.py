# -*- coding: utf-8 -*-
import openpyxl, copy
from openpyxl.styles import PatternFill, Font

f = r"C:\Users\ZQQ\OneDrive\Desktop\IDE\RCA-graduate-diploma U3-U4 project\padlet_card_tags.xlsx"
wb = openpyxl.load_workbook(f)
ws = wb["卡片标签表 Card Tags"]

# capture nav-row style template from current AI Life nav (row 30), cols 1..8
navtpl = [(copy.copy(ws.cell(row=30, column=c).fill),
           copy.copy(ws.cell(row=30, column=c).font),
           copy.copy(ws.cell(row=30, column=c).alignment)) for c in range(1, 9)]

COL = {"QUESTION": "DBE9FF", "INSIGHT": "FFF2CC", "PRECEDENT": "EAD9F5",
       "MAKING": "D9EAD3", "CONTEXT": "EDEDED", "PRIMARY": "FCE4CC"}

def fill(t):
    return PatternFill(start_color=COL[t], end_color=COL[t], fill_type="solid")

N = "③ AI Life 主线"
W = "③ AI Life 野生型"
FI = "③ AI Life 田野访谈"

# (section, type, id, status, date, title, note, is_nav)
rows = [
 (N, None, None, None, None, "▌DIRECTION 1 — AI LIFE", None, True),
 (N, "CONTEXT", "AL-S1", "SUPPORTING", "待填", "SUMMARY — AI 'Living' Entity Project Summary  [AGENT_SUMMARY.pdf]", "claude 起草、Zixun 出想法", False),
 (N, "QUESTION", "AL-Q1", "LOAD-BEARING", "待填", "Can you keep an AI 'alive' on a local machine, like a pet?", "主链起点", False),
 (N, "INSIGHT", "AL-Q2", "LOAD-BEARING", "待填", "Why doesn't an AI think when no one is talking to it?", "问题主链，但内容是推理→黄", False),
 (N, "PRECEDENT", "AL-C1", "SUPPORTING", "待填", "ANSWER 1 | DMN (Raichle et al., 2001)", "", False),
 (N, "PRECEDENT", "AL-C2", "SUPPORTING", "待填", "ANSWER 2 | Symbol Grounding (Harnad 1990; Bender & Koller 2020)", "", False),
 (N, "INSIGHT", "AL-Q3", "LOAD-BEARING", "待填", "What if I wire it up to sensors?", "问题主链，但内容是推理→黄", False),
 (N, "INSIGHT", "AL-Q4", "LOAD-BEARING", "待填", "Are the three problems separate, or does one cause the others?", "问题主链，但内容是推理→黄", False),
 (N, "INSIGHT", "AL-Q5", "LOAD-BEARING", "待填", "What would the smallest system that runs all three look like?", "问题主链，但内容是设计论证→黄", False),
 (N, "INSIGHT", "AL-Q6", "LOAD-BEARING", "待填", "When the model 'remembers,' what has actually changed?", "问题主链，但内容是推理→黄", False),
 (N, "PRECEDENT", "AL-R0", "SUPPORTING", "待填", "From Text to Life (Nisioti et al., 2024)", "", False),
 (N, "INSIGHT", "AL-R1", "LOAD-BEARING", "待填", "Why I went looking for this paper at all", "论文精读线，但内容是反思→黄", False),
 (N, "INSIGHT", "AL-R2", "LOAD-BEARING", "待填", "The paper gave me a definition I could stand on", "论文精读线，但内容是反思→黄", False),
 (N, "INSIGHT", "AL-R3", "LOAD-BEARING", "待填", "Does an LLM have internal states?", "论文精读线，但内容是反思→黄", False),
 (N, "PRECEDENT", "AL-R4", "SUPPORTING", "待填", "R4) ELM: what 'improvement' really means", "", False),
 (N, "PRECEDENT", "AL-R5", "SUPPORTING", "待填", "R5) LMX: the same move, domain-independent", "", False),
 (N, "INSIGHT", "AL-R6", "LOAD-BEARING", "待填", "R6) Learning lives in the population, not the weights", "论文精读线，但内容是改变想法→黄", False),
 (N, "INSIGHT", "AL-R7", "LOAD-BEARING", "待填", "R7) A sharper question, and one I had to throw away", "论文精读线，但内容是反思→黄", False),
 (N, "QUESTION", "AL-Q7", "LOAD-BEARING", "待填", "Does it develop behaviour that was never designed into it?", "", False),
 (N, "INSIGHT", "AL-I1", "LOAD-BEARING", "待填", "Why I replaced the GRU emotion core with CfC + AutoNCP", "原误标 AL-L1 → 改 I 前缀", False),
 (N, "PRECEDENT", "AL-S3", "SUPPORTING", "待填", "AI Life-Form Architecture Validation: Neuroscience Lit Review  [agent_en.pdf]", "claude 撰写；S=文档角色", False),
 (N, "CONTEXT", "AL-S2", "SUPPORTING", "待填", "Project phase divided  [AGENT_PHASES.pdf]", "claude 起草的阶段规划", False),
 (N, "INSIGHT", "AL-I2", "SUPPORTING", "待填", "Step 1: Design Concept & Technology Selection  [phase_en.pdf]", "设计/选型决策，非动手做", False),
 (N, "MAKING", "AL-M1", "RUNNING", "待填", "Built a Python AI-life simulator (build overview)", "", False),
 (N, "MAKING", "AL-M2", "RUNNING", "待填", "Step 2: Phase 1 — Building the Core Emotion Network  [life_v2.py]", "附件 life_v2.py", False),
 (N, "MAKING", "AL-M3", "RUNNING", "待填", "Step 3: Phase 2 — Layering Biological Systems", "", False),
 (N, "MAKING", "AL-M4", "SUPPORTING", "待填", "Supplementary: Pre-Integration Smoke Test  [test_world.py]", "", False),
 (N, "MAKING", "AL-M5", "RUNNING", "待填", "Step 4: World Simulator  [world_v2.py]", "附件 world_v2.py", False),
 (N, "MAKING", "AL-M6", "RUNNING", "待填", "Step 5: Discovering & Fixing Three Training Bugs", "", False),
 (N, "MAKING", "AL-M7", "RUNNING", "待填", "Step 6: Phase 3/4 Interface Placeholders", "", False),
 (N, "MAKING", "AL-M8", "SUPPORTING", "待填", "Step 7: Scheduled Task — Automated Code Injection  [phase_status.json]", "附件 phase_status.json", False),
 (N, "MAKING", "AL-M9", "SUPPORTING", "待填", "Step 8: Real-time Monitoring Dashboard  [dashboard.py]", "附件 dashboard.py", False),
 (N, "MAKING", "AL-M10", "SUPPORTING", "待填", "Dashboard: Full Bilingual UI", "", False),
 (N, "MAKING", "AL-M11", "RUNNING", "待填", "Step 9: Switching to a Local LLM via Ollama (qwen3:14b)", "", False),
 (N, "INSIGHT", "AL-I3", "RUNNING", "2026.6.15", "Developmental stage dev=1.0 — full training analysis + arousal-freeze anomaly  [life_log.jsonl]", "附件 life_log.jsonl", False),
 (N, "INSIGHT", "AL-I4", "LOAD-BEARING", "待填", "Why I Made This Change (worldbook / missing injection / clean restart)  [worldbook.py]", "附件 worldbook.py", False),
 (N, "MAKING", "AL-M12", "RUNNING", "待填", "step 10: Discover the Missing Injection  [world_v2.py]", "", False),
 (N, "MAKING", "AL-M13", "RUNNING", "待填", "Step 11: Reconnect the World Setting", "原无 ID → 补", False),
 (N, "MAKING", "AL-M14", "RUNNING", "待填", "Step 12: Expand the Context Window (num_ctx 16384)", "原无 ID → 补", False),
 (N, "MAKING", "AL-M15", "RUNNING", "待填", "Step 13: Reinforce the Bilingual Rule", "原无 ID → 补", False),
 (N, "MAKING", "AL-M16", "RUNNING", "待填", "Step 14: Back Up and Restart Clean & Verify  [emotion_weights_v2.pt]", "附件 emotion_weights_v2.pt；原无 ID → 补", False),
 (N, "CONTEXT", "AL-F5", "SUPPORTING", "待填", "__A_B__.md  (实验A & 实验B 设计总纲)", "文件独立帖；删 .pyc、删重复 life_log.jsonl", False),
 (W, None, None, None, None, "▌DIRECTION 1 — AI LIFE · 野生型 Plan B", None, True),
 (W, "INSIGHT", "AL-I5", "LOAD-BEARING", "待填", "Experiment A · Wild Type Architecture — Core Decisions", "原标题写 Experiment B → 改 A", False),
 (W, "INSIGHT", "AL-I6", "LOAD-BEARING", "待填", "Step1: CfC Neural Core — 22D Input vs 36D  [life_wild.py]", "附件 life_wild.py；控制变量论证", False),
 (W, "INSIGHT", "AL-I7", "LOAD-BEARING", "待填", "Step2: Removed Modules — Stripping All Language & Narrative", "删除决策论证", False),
 (W, "MAKING", "AL-M17", "RUNNING", "待填", "Step3: New Modules — Sensor Bridge + Internal Symbol Tracker  [sensor_bridge.py]", "真新建模块；附件 sensor_bridge.py", False),
 (W, "INSIGHT", "AL-I8", "LOAD-BEARING", "待填", "Step4: Preserved Core Mechanisms — Comparable Baseline", "保留理由论证", False),
 (W, "INSIGHT", "AL-I9", "LOAD-BEARING", "待填", "Perception Source — Arduino Physical Sensors Replace LLM World", "感知层总述（读起来建议放 Step1 前）", False),
 (W, "MAKING", "AL-M18", "LOAD-BEARING", "待填", "symbol_tracker.py  (野生型核心观测仪器)", "文件独立帖；K-Means+转移矩阵+符号稳定度", False),
 (FI, None, None, None, None, "▌FIELD INVESTIGATE — AI Life 田野访谈", None, True),
 (FI, "PRIMARY", "AL-P0", "SUPPORTING", "待填", "field_interview_compact.docx  [访谈问卷模板]", "一手研究工具", False),
 (FI, "PRIMARY", "AL-P1", "LOAD-BEARING", "待填", "FIELD INVESTIGATE — Suwen Yuan  [suwen_yuan.wps]", "", False),
 (FI, "PRIMARY", "AL-P2", "LOAD-BEARING", "待填", "FIELD INVESTIGATE — Yuying Zhang  [yuying_zhang.docx]", "", False),
 (FI, "PRIMARY", "AL-P3", "LOAD-BEARING", "待填", "FIELD INVESTIGATE — Qiqi  [qiqi.docx]", "", False),
 (FI, "PRIMARY", "AL-P4", "LOAD-BEARING", "待填", "FIELD INVESTIGATE — Z  [Z.docx]", "", False),
]

print("NEW ROW COUNT", len(rows))
ws.delete_rows(30, 71)
ws.insert_rows(30, len(rows))
for i, (sec, typ, rid, st, dt, title, note, nav) in enumerate(rows):
    r = 30 + i
    ws.cell(row=r, column=2, value=sec)
    ws.cell(row=r, column=7, value=title)
    if nav:
        for c in range(1, 9):
            fl, fo, al = navtpl[c - 1]
            ws.cell(row=r, column=c).fill = copy.copy(fl)
            ws.cell(row=r, column=c).font = copy.copy(fo)
            ws.cell(row=r, column=c).alignment = copy.copy(al)
        ws.cell(row=r, column=7, value=title)
    else:
        ws.cell(row=r, column=3, value=typ)
        ws.cell(row=r, column=3).fill = fill(typ)
        ws.cell(row=r, column=3).font = Font(bold=True)
        ws.cell(row=r, column=4, value=rid)
        ws.cell(row=r, column=5, value=st)
        ws.cell(row=r, column=6, value=dt)
        ws.cell(row=r, column=8, value=note)

# append PR-08 personal reflection (moved out of AI Life)
last = ws.max_row + 1
ws.cell(row=last, column=2, value="④ 作品集反思")
ws.cell(row=last, column=3, value="INSIGHT")
ws.cell(row=last, column=3).fill = fill("INSIGHT")
ws.cell(row=last, column=3).font = Font(bold=True)
ws.cell(row=last, column=4, value="PR-08")
ws.cell(row=last, column=5, value="SUPPORTING")
ws.cell(row=last, column=6, value="待填")
ws.cell(row=last, column=7, value="personal reflection — Western portfolios & white space")
ws.cell(row=last, column=8, value="原夹在 AI Life 里，移到此板块")

# renumber the '#' column
cnt = 0
for r in ws.iter_rows(min_row=2):
    cnt += 1
    r[0].value = cnt

try:
    wb.save(f)
    print("SAVED OK. new max_row", ws.max_row)
except PermissionError:
    print("PERMISSION_ERROR: 文件被 Excel 占用，请关闭后重试")
